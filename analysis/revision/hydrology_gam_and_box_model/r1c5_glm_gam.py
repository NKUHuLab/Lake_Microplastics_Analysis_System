# -*- coding: utf-8 -*-
"""R1C5 Fig. 2 rebuild: original visuals, revised GLM/GAM internals.

This script intentionally preserves the visual grammar of the manuscript code:
- 18_advanced_meta_analysis.py: dumbbell plot with black rods and three colored dots.
- 19_quantify_effects.py: faceted predictor-response plots stratified by moderator.

The statistical internals are revised for the reviewer request:
- subgroup and global moderation coefficients use Gaussian GLM on ln(MP) with
  HC3 robust standard errors, 95% CIs and p values;
- linear response lines are replaced by GAM smooths with 95% CIs.
"""

from __future__ import annotations

from pathlib import Path
import os
import io
import math
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

import numpy as np
import pandas as pd
import matplotlib as mpl
import matplotlib.pyplot as plt
from sklearn.decomposition import PCA
from sklearn.preprocessing import StandardScaler

REPO_ROOT = Path(os.environ.get("LAKE_MP_REPO_ROOT", Path(__file__).resolve().parents[3]))
DATA_ROOT = Path(os.environ.get("LAKE_MP_DATA_ROOT", REPO_ROOT / "data"))
CODE_DIR = Path(os.environ.get("LAKE_MP_GAM_HELPERS_DIR", Path(__file__).resolve().parent))
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

from r1c5_glm_gam_helpers import (  # noqa: E402
    PREDICTORS,
    MODERATORS,
    fit_subgroup_glm_hc3,
    fit_global_moderation_glm_hc3,
    gam_fit_curve,
    make_tertiles,
    vif_table,
)


BASE = Path(__file__).resolve().parents[2]
TRAIN = Path(os.environ.get("LAKE_MP_TRAIN_DATA", DATA_ROOT / "model_products" / "train_data.csv"))
OUT_ROOT = Path(os.environ.get("LAKE_MP_OUTPUT_DIR", REPO_ROOT / "outputs")) / "hydrology_gam_and_box_model"
BASE = OUT_ROOT
OUT_FIG = BASE / "re_fig" / "R1_5_GLMM调节效应"
OUT_DATA = OUT_ROOT / "data"
OUT_FIG.mkdir(parents=True, exist_ok=True)
OUT_DATA.mkdir(parents=True, exist_ok=True)


mpl.rcParams.update({
    "font.family": "sans-serif",
    "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
    "pdf.fonttype": 42,
    "ps.fonttype": 42,
    "svg.fonttype": "none",
    "axes.unicode_minus": False,
    "font.size": 12,
    "axes.spines.top": False,
    "axes.spines.right": False,
})

COLORS_ORIGINAL = ["#1f77b4", "#ff7f0e", "#2ca02c"]
COLORS_V3 = ["#466A77", "#5B7352", "#DA3839"]
GROUPS = ["Low", "Medium", "High"]
RT_GROUPS = ["Low RT", "Medium RT", "High RT"]
DUMBBELL_PLOT_MODERATORS = {"Res_time", "Shore_dev"}

MOD_LABELS = {
    "Res_time": "Residence time",
    "Shore_dev": "Shoreline development",
    "Lake_area": "Lake area",
    "Vol_total": "Total volume",
}


def save_pub(fig: plt.Figure, stem: str, dpi: int = 300) -> None:
    for ext in ("png", "pdf"):
        fig.savefig(
            OUT_FIG / f"{stem}.{ext}",
            dpi=dpi,
            bbox_inches="tight",
            metadata={"Creator": "Codex matplotlib; editable text PDF"},
        )
    plt.close(fig)


def load_data() -> pd.DataFrame:
    df = pd.read_csv(TRAIN)
    df = df.dropna(subset=["ln"]).copy()
    for col in PREDICTORS + MODERATORS:
        if col not in df.columns:
            df[col] = 0.0
    return df


def plot_dumbbell_original_style(
    effects_df: pd.DataFrame,
    group_by_var: str,
    ci_low_df: pd.DataFrame | None = None,
    ci_high_df: pd.DataFrame | None = None,
    feature_order: list[str] | None = None,
    sort_ascending: bool = False,
    suffix: str = "v4",
) -> None:
    """Visual style matched to 18_advanced_meta_analysis.py."""
    effects_df = effects_df.copy()
    if feature_order is None:
        effects_df["abs_mean"] = effects_df.abs().mean(axis=1)
        df_sorted = effects_df.sort_values("abs_mean", ascending=sort_ascending).drop(columns="abs_mean")
    else:
        ordered = [f for f in feature_order if f in effects_df.index]
        ordered += [f for f in effects_df.index if f not in ordered]
        df_sorted = effects_df.loc[ordered]
    ci_low_sorted = ci_low_df.loc[df_sorted.index] if ci_low_df is not None else None
    ci_high_sorted = ci_high_df.loc[df_sorted.index] if ci_high_df is not None else None

    group_names = list(df_sorted.columns)
    y_pos = np.arange(len(df_sorted))
    fig, ax = plt.subplots(figsize=(14, 12))

    ax.hlines(
        y=y_pos,
        xmin=df_sorted.min(axis=1),
        xmax=df_sorted.max(axis=1),
        color="black",
        alpha=0.7,
        linewidth=2.5,
        zorder=1,
    )

    for i, group in enumerate(group_names):
        color = COLORS_ORIGINAL[i % len(COLORS_ORIGINAL)]
        if ci_low_sorted is not None and ci_high_sorted is not None:
            coef = df_sorted[group].to_numpy(float)
            xerr = np.vstack([
                coef - ci_low_sorted[group].to_numpy(float),
                ci_high_sorted[group].to_numpy(float) - coef,
            ])
            ax.errorbar(
                coef,
                y_pos,
                xerr=xerr,
                fmt="o",
                color=color,
                markersize=12,
                linewidth=1.5,
                capsize=4,
                label=group,
                zorder=3,
                markeredgecolor="black",
                markeredgewidth=0.5,
            )
        else:
            ax.scatter(
                df_sorted[group],
                y_pos,
                color=color,
                s=180,
                label=group,
                zorder=3,
                edgecolors="black",
                linewidth=0.5,
            )

    xmin, xmax = ax.get_xlim()
    ax.axvspan(xmin, 0, color="royalblue", alpha=0.05, zorder=0, label="_nolegend_")
    ax.axvspan(0, xmax, color="tomato", alpha=0.05, zorder=0, label="_nolegend_")

    ax.xaxis.grid(True, linestyle="--", which="major", color="grey", alpha=0.3)
    ax.yaxis.grid(False)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.tick_params(axis="y", length=0)

    ax.set_xlabel("Effect coefficient (standardized ln scale; 95% CI)", fontsize=14)
    ax.set_ylabel("")
    ax.set_title(
        f"Moderation Effect by {MOD_LABELS.get(group_by_var, group_by_var)}",
        fontsize=22,
        pad=20,
    )
    legend = ax.legend(
        title=f"Stratified by {MOD_LABELS.get(group_by_var, group_by_var)}",
        fontsize=12,
        loc="upper right",
        bbox_to_anchor=(1.15, 0.95),
    )
    plt.setp(legend.get_title(), fontsize=14)
    ax.tick_params(axis="x", labelsize=12)
    ax.set_yticks(y_pos)
    ax.set_yticklabels(df_sorted.index, fontdict={"family": "Arial", "size": 12})
    plt.tight_layout(rect=[0, 0, 0.9, 1])
    save_pub(fig, f"Fig_R1_5_Dumbbell_{group_by_var}_GLM_originalStyle_{suffix}")


def plot_dumbbell_focus_ci(
    effects_df: pd.DataFrame,
    group_by_var: str,
    ci_low_df: pd.DataFrame,
    ci_high_df: pd.DataFrame,
    xlim: tuple[float, float] = (-2.2, 2.2),
    feature_order: list[str] | None = None,
    sort_ascending: bool = False,
    compact_ci_pairs: set[tuple[str, str]] | None = None,
    compact_ci_min_width: float = 2.5,
    compact_ci_half_width: float = 0.42,
    figsize: tuple[float, float] = (13.09, 11.6),
    suffix: str = "focusCI_v4",
) -> None:
    """Main-text friendly dumbbell plot with off-scale CI arrows.

    Full CI values remain in the table/CSV. The figure keeps the original
    manuscript visual grammar but prevents one unstable interval from flattening
    the rest of the coefficients.
    """
    effects_df = effects_df.copy()
    if feature_order is None:
        effects_df["abs_mean"] = effects_df.abs().mean(axis=1)
        df_sorted = effects_df.sort_values("abs_mean", ascending=sort_ascending).drop(columns="abs_mean")
    else:
        ordered = [f for f in feature_order if f in effects_df.index]
        ordered += [f for f in effects_df.index if f not in ordered]
        df_sorted = effects_df.loc[ordered]
    ci_low_sorted = ci_low_df.loc[df_sorted.index]
    ci_high_sorted = ci_high_df.loc[df_sorted.index]
    y_pos = np.arange(len(df_sorted))
    group_names = list(df_sorted.columns)
    xmin, xmax = xlim

    fig, ax = plt.subplots(figsize=figsize)
    ax.hlines(
        y=y_pos,
        xmin=df_sorted.min(axis=1),
        xmax=df_sorted.max(axis=1),
        color="black",
        alpha=0.7,
        linewidth=2.5,
        zorder=1,
    )

    any_truncated = False
    any_compacted = False
    cap_half_height = 0.045
    for i, group in enumerate(group_names):
        color = COLORS_ORIGINAL[i % len(COLORS_ORIGINAL)]
        coefs = df_sorted[group].to_numpy(float)
        lows = ci_low_sorted[group].to_numpy(float)
        highs = ci_high_sorted[group].to_numpy(float)
        for feature, y, coef, lo, hi in zip(df_sorted.index, y_pos, coefs, lows, highs):
            compact_this_ci = (
                compact_ci_pairs is not None
                and (feature, group) in compact_ci_pairs
                and (hi - lo) >= compact_ci_min_width
            )
            if compact_this_ci:
                any_compacted = True
                lo_compact = max(coef - compact_ci_half_width, xmin)
                hi_compact = min(coef + compact_ci_half_width, xmax)
                ax.plot(
                    [lo_compact, hi_compact],
                    [y, y],
                    color=color,
                    lw=1.35,
                    alpha=0.72,
                    linestyle=(0, (2.0, 1.7)),
                    zorder=2,
                )
                ax.scatter(lo_compact, y, marker="<", s=42, color=color, edgecolor="none", alpha=0.9, zorder=2.6)
                ax.scatter(hi_compact, y, marker=">", s=42, color=color, edgecolor="none", alpha=0.9, zorder=2.6)
                continue
            lo_clip = max(lo, xmin)
            hi_clip = min(hi, xmax)
            ax.plot([lo_clip, hi_clip], [y, y], color=color, lw=1.45, alpha=0.92, zorder=2)
            if lo >= xmin:
                ax.plot([lo, lo], [y - cap_half_height, y + cap_half_height], color=color, lw=1.15, zorder=2)
            else:
                any_truncated = True
                ax.scatter(xmin, y, marker="<", s=42, color=color, edgecolor="none", zorder=2.6)
            if hi <= xmax:
                ax.plot([hi, hi], [y - cap_half_height, y + cap_half_height], color=color, lw=1.15, zorder=2)
            else:
                any_truncated = True
                ax.scatter(xmax, y, marker=">", s=42, color=color, edgecolor="none", zorder=2.6)
        ax.scatter(
            coefs,
            y_pos,
            color=color,
            s=180,
            label=group,
            zorder=3,
            edgecolors="black",
            linewidth=0.5,
        )

    ax.set_xlim(xmin, xmax)
    ax.axvspan(xmin, 0, color="royalblue", alpha=0.05, zorder=0, label="_nolegend_")
    ax.axvspan(0, xmax, color="tomato", alpha=0.05, zorder=0, label="_nolegend_")
    ax.xaxis.grid(True, linestyle="--", which="major", color="grey", alpha=0.3)
    ax.yaxis.grid(False)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.set_xlabel("Effect coefficient (standardized ln scale; 95% CI)", fontsize=14)
    ax.set_ylabel("")
    ax.set_title(
        f"Moderation Effect by {MOD_LABELS.get(group_by_var, group_by_var)}",
        fontsize=19,
        pad=20,
    )
    legend = ax.legend(
        title=f"Stratified by {MOD_LABELS.get(group_by_var, group_by_var)}",
        fontsize=9.5,
        loc="upper left",
        bbox_to_anchor=(1.015, 0.94),
        ncol=1,
        columnspacing=0.8,
        handletextpad=0.45,
    )
    plt.setp(legend.get_title(), fontsize=10.5)
    ax.tick_params(axis="x", labelsize=12)
    ax.set_yticks(y_pos)
    ax.set_yticklabels(df_sorted.index, fontdict={"family": "Arial", "size": 12})
    note_lines = []
    if any_truncated:
        note_lines.append("Edge arrowheads: 95% CIs exceed the x-window.")
    if any_compacted:
        note_lines.append("Paired chevrons: compacted very wide Primary_Waste_Discharge Low-group CIs.")
        note_lines.append("Exact interval limits are reported in the source table.")
    if note_lines:
        ax.text(
            0.5,
            -0.115,
            "\n".join(note_lines),
            transform=ax.transAxes,
            fontsize=8.5,
            color="#666666",
            ha="center",
            va="top",
        )
    plt.tight_layout(rect=[0, 0.04, 0.82, 1])
    save_pub(fig, f"Fig_R1_5_Dumbbell_{group_by_var}_GLM_originalStyle_{suffix}")


def make_dumbbell_outputs(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    all_subgroup_stats = []
    all_global_stats = []
    all_diag = []
    all_vif = []

    for moderator in MODERATORS:
        df[f"{moderator}_group"] = make_tertiles(df, moderator, GROUPS)
        group_effects = {}
        group_ci_low = {}
        group_ci_high = {}
        for group in GROUPS:
            sub = df[df[f"{moderator}_group"] == group]
            coefs, stats_df, diag = fit_subgroup_glm_hc3(sub, PREDICTORS, group)
            if coefs is None or stats_df is None:
                continue
            group_effects[group] = coefs
            stats_indexed = stats_df[stats_df["Feature"] != "Intercept"].set_index("Feature")
            group_ci_low[group] = stats_indexed["CI_Lower_95"]
            group_ci_high[group] = stats_indexed["CI_Upper_95"]
            stats_df.insert(1, "Moderator", moderator)
            all_subgroup_stats.append(stats_df)
            if diag is not None:
                all_diag.append({"model_scope": "subgroup", "moderator": moderator, "group": group, **diag})

        if group_effects:
            comparison_df = pd.DataFrame(group_effects)
            ci_low_df = pd.DataFrame(group_ci_low).loc[comparison_df.index]
            ci_high_df = pd.DataFrame(group_ci_high).loc[comparison_df.index]
            max_abs_order_table = (
                comparison_df.abs()
                .max(axis=1)
                .rename("max_abs_effect")
                .sort_values(ascending=True)
                .reset_index()
                .rename(columns={"index": "Feature"})
            )
            max_abs_order_table["plot_rank_ascending"] = np.arange(1, len(max_abs_order_table) + 1)
            max_abs_order_table.to_csv(
                OUT_DATA / f"R1C5_GLM_originalStyle_dumbbell_maxAbs_order_{moderator}.csv",
                index=False,
                encoding="utf-8-sig",
            )
            max_abs_order = max_abs_order_table["Feature"].tolist()
            comparison_df.to_csv(
                OUT_DATA / f"R1C5_GLM_originalStyle_dumbbell_coefficients_{moderator}.csv",
                encoding="utf-8-sig",
            )
            if moderator in DUMBBELL_PLOT_MODERATORS:
                plot_dumbbell_focus_ci(
                    comparison_df,
                    moderator,
                    ci_low_df,
                    ci_high_df,
                    feature_order=max_abs_order,
                    compact_ci_pairs={("Primary_Waste_Discharge", "Low")},
                    suffix="compactPrimaryLowCI_maxAbsAscending_v4",
                )

        global_df, global_diag = fit_global_moderation_glm_hc3(df, moderator, PREDICTORS)
        all_global_stats.append(global_df)
        all_diag.append({"model_scope": "global", "moderator": moderator, "group": "all", **global_diag})
        all_vif.append(vif_table(df, PREDICTORS + [moderator]).assign(Moderator=moderator))

    subgroup_out = pd.concat(all_subgroup_stats, ignore_index=True)
    global_out = pd.concat(all_global_stats, ignore_index=True)
    diag_out = pd.DataFrame(all_diag)
    vif_out = pd.concat(all_vif, ignore_index=True)

    subgroup_out.to_csv(OUT_DATA / "R1C5_GLM_subgroup_coefficients_originalStyle_v4.csv", index=False, encoding="utf-8-sig")
    global_out.to_csv(OUT_DATA / "R1C5_GLM_global_moderation_coefficients_originalStyle_v4.csv", index=False, encoding="utf-8-sig")
    diag_out.to_csv(OUT_DATA / "R1C5_GLM_model_diagnostics_originalStyle_v4.csv", index=False, encoding="utf-8-sig")
    vif_out.to_csv(OUT_DATA / "R1C5_GLM_vif_originalStyle_v4.csv", index=False, encoding="utf-8-sig")
    return global_out, diag_out


def make_gam_fishery_rt(df: pd.DataFrame) -> pd.DataFrame:
    fig, axes = plt.subplots(1, 3, figsize=(13, 4.5))
    df = df.copy()
    df["_rtg"] = make_tertiles(df, "Res_time", RT_GROUPS)
    rows = []

    for i, (gname, ax) in enumerate(zip(RT_GROUPS, axes)):
        sub = df[df["_rtg"] == gname].copy()
        x = np.log1p(sub["fish_gdp_sqkm"].clip(lower=0).to_numpy(float))
        y = sub["ln"].to_numpy(float)
        fit = gam_fit_curve(x, y, n_splines=6, grid_size=200)
        rows.append({
            "panel": gname,
            "N": fit["n"],
            "GAM_pseudo_R2": fit["pseudo_r2"],
            "Linear_R2_reference": fit["linear_r2"],
        })

        ax.scatter(x, y, s=12, alpha=0.35, color=COLORS_V3[i], edgecolors="white", linewidth=0.15, zorder=1)
        ax.plot(fit["x"], fit["fit"], color=COLORS_V3[i], lw=2, zorder=3)
        ax.fill_between(fit["x"], fit["ci_low"], fit["ci_high"], color=COLORS_V3[i], alpha=0.12, zorder=2)
        ax.set_xlabel("log1p(Fishery GDP)", fontsize=8.5)
        ax.set_ylabel("ln(MP abundance)", fontsize=8.5)
        ax.set_title(f"{chr(65 + i)}. {gname} (N={fit['n']})", fontsize=9, fontweight="bold", loc="left")
        ax.text(
            0.03,
            0.95,
            f"GAM pseudo-R2={fit['pseudo_r2']:.2f}\nLinear R2={fit['linear_r2']:.2f}",
            transform=ax.transAxes,
            fontsize=7,
            va="top",
        )

    fig.suptitle("Nonlinear: Fishery GDP vs MP Abundance (GAM, 95% CI)", fontsize=10.5, fontweight="bold", y=1.02)
    plt.tight_layout()
    save_pub(fig, "Fig_R1_5_GAM_Nonlinear_v4")
    out = pd.DataFrame(rows)
    out.to_csv(OUT_DATA / "R1C5_GAM_fishery_RT_model_summary_v4.csv", index=False, encoding="utf-8-sig")
    return out


def create_wastewater_index(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    waste_components = ["Primary_Waste_Discharge", "Secondary_Waste_Discharge", "Advanced_Waste_Discharge"]
    X = StandardScaler().fit_transform(out[waste_components].fillna(0))
    out["Wastewater_PC1"] = PCA(n_components=1).fit_transform(X)
    corr = np.corrcoef(out["Wastewater_PC1"], out["Advanced_Waste_Discharge"].fillna(0))[0, 1]
    if np.isfinite(corr) and corr < 0:
        out["Wastewater_PC1"] = -out["Wastewater_PC1"]
    out["fish_gdp_sqkm (log-transformed)"] = np.log1p(out["fish_gdp_sqkm"].clip(lower=0))
    return out


def make_gam_interaction_grids(df: pd.DataFrame) -> None:
    """Replacement for 19_quantify_effects.py lmplot: same facets, nonlinear GAM."""
    df = create_wastewater_index(df)
    key_predictors = ["Wastewater_PC1", "Cultivated_land", "fish_gdp_sqkm (log-transformed)"]

    for mod_var in ["Res_time", "Shore_dev"]:
        group_col = f"{mod_var}_group"
        df[group_col] = make_tertiles(df, mod_var, GROUPS)
        fig, axes = plt.subplots(1, 3, figsize=(19.8, 6), sharey=False)
        for ax, predictor in zip(axes, key_predictors):
            for i, group in enumerate(GROUPS):
                sub = df[df[group_col] == group].dropna(subset=[predictor, "ln"])
                x = sub[predictor].to_numpy(float)
                y = sub["ln"].to_numpy(float)
                ax.scatter(x, y, alpha=0.4, s=50, color=COLORS_ORIGINAL[i], label=group, zorder=1)
                try:
                    fit = gam_fit_curve(x, y, n_splines=6, grid_size=160)
                    ax.plot(fit["x"], fit["fit"], color=COLORS_ORIGINAL[i], lw=2.0, zorder=3)
                    ax.fill_between(fit["x"], fit["ci_low"], fit["ci_high"], color=COLORS_ORIGINAL[i], alpha=0.10, zorder=2)
                except Exception:
                    continue
            ax.set_title(predictor, fontsize=16)
            ax.set_xlabel("Predictor value", fontsize=13)
            ax.set_ylabel("ln(Microplastic abundance)", fontsize=13)
            ax.grid(True, linestyle="--", color="grey", alpha=0.25)

        handles, labels = axes[-1].get_legend_handles_labels()
        fig.legend(handles[:3], labels[:3], title=f"{mod_var} level", loc="upper right", bbox_to_anchor=(0.98, 0.92))
        fig.suptitle(f"Interaction Effects Moderated by {MOD_LABELS.get(mod_var, mod_var)} (GAM smooth)", y=1.03, fontsize=22)
        plt.tight_layout(rect=[0, 0, 0.95, 1])
        save_pub(fig, f"Fig_R1_5_19_GAM_interaction_by_{mod_var}_v4")


def make_three_line_docx(global_out: pd.DataFrame, diag_out: pd.DataFrame) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    out = OUT_DATA / "R1C5_GLM_moderation_three_line_table_originalStyle_v4.docx"
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(12)

    def add_p(text: str, bold: bool = False) -> None:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(12)

    def borders(cell, top=None, bottom=None):
        tc_pr = cell._tc.get_or_add_tcPr()
        b = tc_pr.first_child_found_in("w:tcBorders")
        if b is None:
            b = OxmlElement("w:tcBorders")
            tc_pr.append(b)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = b.find(qn(f"w:{edge}"))
            if el is None:
                el = OxmlElement(f"w:{edge}")
                b.append(el)
            if edge == "top" and top:
                el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(top)); el.set(qn("w:color"), "000000")
            elif edge == "bottom" and bottom:
                el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(bottom)); el.set(qn("w:color"), "000000")
            else:
                el.set(qn("w:val"), "none")

    def cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = align
        r = p.add_run(str(text))
        r.bold = bold
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(12)

    add_p("Table R1C5-S2. Moderation model output for residence-time amplification.", True)
    keep_terms = [
        "fish_gdp_sqkm",
        "Res_time",
        "Shore_dev",
        "fish_gdp_sqkm:Res_time",
        "fish_gdp_sqkm:Shore_dev",
    ]
    sub = global_out[global_out["Feature"].isin(keep_terms)].copy()
    sub = sub[sub["Context"].isin(["Res_time", "Shore_dev"])]
    headers = ["Moderator", "Term", "Estimate", "SE", "95% CI", "p", "Effect (%)"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, bold=True)
        borders(table.rows[0].cells[i], top=16, bottom=8)
    for _, r in sub.iterrows():
        row = table.add_row().cells
        vals = [
            MOD_LABELS.get(r["Context"], r["Context"]),
            r["Feature"],
            f"{r['Coefficient']:.3f}",
            f"{r['Std_Error']:.3f}",
            f"{r['CI_Lower_95']:.3f} to {r['CI_Upper_95']:.3f}",
            f"{r['P_value']:.3g}",
            f"{r['Effect_percent']:.1f}",
        ]
        for i, val in enumerate(vals):
            cell_text(row[i], val, align=WD_ALIGN_PARAGRAPH.LEFT if i in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER)
            borders(row[i])
    for cell in table.rows[-1].cells:
        borders(cell, bottom=16)

    add_p("Note: The primary inferential model is a Gaussian GLM on ln(MP abundance), matching the original response scale. HC3 robust standard errors are used for 95% confidence intervals and p values. A GLMM was not used because the training table does not contain a study, country, campaign, lake, or site identifier for a defensible random-effect structure.")
    add_p("Residence-time source: HydroLAKES v1.0 hydraulic residence time, calculated as lake volume divided by outflow (RT = V/Qout), was used as the hydrological moderator.")

    add_p("Table R1C5-S3. Diagnostics for global moderation models.", True)
    dtable = doc.add_table(rows=1, cols=6)
    dtable.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Moderator", "N", "AIC", "R2", "RMSE", "df"]):
        cell_text(dtable.rows[0].cells[i], h, bold=True)
        borders(dtable.rows[0].cells[i], top=16, bottom=8)
    for _, r in diag_out[diag_out["model_scope"] == "global"].iterrows():
        row = dtable.add_row().cells
        vals = [MOD_LABELS.get(r["moderator"], r["moderator"]), int(r["N"]), f"{r['AIC']:.1f}", f"{r['R2']:.3f}", f"{r['RMSE']:.2f}", int(r["df_resid"])]
        for i, val in enumerate(vals):
            cell_text(row[i], val)
            borders(row[i])
    for cell in dtable.rows[-1].cells:
        borders(cell, bottom=16)
    doc.save(out)
    print(f"Saved DOCX: {out}")


def main() -> None:
    df = load_data()
    print(f"Loaded {len(df)} observations from {TRAIN}")
    global_out, diag_out = make_dumbbell_outputs(df)
    gam_summary = make_gam_fishery_rt(df)
    make_gam_interaction_grids(df)
    make_three_line_docx(global_out, diag_out)
    print(gam_summary.to_string(index=False))
    print(f"Saved figures to: {OUT_FIG}")
    print(f"Saved data/tables to: {OUT_DATA}")


if __name__ == "__main__":
    main()
