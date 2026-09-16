# -*- coding: utf-8 -*-
"""Pre-check Hydrology high-RT evidence using bootstrap/permutation and marginal effects.

This script keeps the same variable transformation as the Hydrology v4 figure code:
Gaussian linear model on ln(MP), log1p for skewed variables, z-standardized
predictors, and HC3 covariance for Wald tests.
"""

from __future__ import annotations

import math
from pathlib import Path
import os

import numpy as np
import pandas as pd


BASE = Path(__file__).resolve().parents[2]
REPO_ROOT = Path(os.environ.get("LAKE_MP_REPO_ROOT", Path(__file__).resolve().parents[3]))
DATA_ROOT = Path(os.environ.get("LAKE_MP_DATA_ROOT", REPO_ROOT / "data"))
TRAIN = Path(os.environ.get("LAKE_MP_TRAIN_DATA", DATA_ROOT / "model_products" / "train_data.csv"))
OUT = Path(os.environ.get("LAKE_MP_OUTPUT_DIR", REPO_ROOT / "outputs")) / "hydrology_gam_and_box_model" / "data"
OUT.mkdir(parents=True, exist_ok=True)

PREDICTORS = [
    "Primary_Waste_Discharge",
    "Secondary_Waste_Discharge",
    "Advanced_Waste_Discharge",
    "RSE_paved",
    "RSE_gravel",
    "RSE_other",
    "Mismanaged",
    "fish_gdp_sqkm",
    "Cultivated_land",
    "Artificial_surface",
]
LOG_TRANSFORM = {"Mismanaged", "fish_gdp_sqkm", "Res_time", "Lake_area", "Vol_total"}
RNG_SEED = 20260707
N_BOOT = 5000
N_PERM = 5000


def norm_cdf(z: float) -> float:
    return 0.5 * (1.0 + math.erf(z / math.sqrt(2.0)))


def norm_sf(z: float) -> float:
    return 1.0 - norm_cdf(z)


def load_data() -> pd.DataFrame:
    df = pd.read_csv(TRAIN)
    needed = ["ln", "Res_time"] + PREDICTORS
    for col in needed:
        if col not in df.columns:
            df[col] = np.nan
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df = df.dropna(subset=needed).copy()
    return df


def transform_matrix(df: pd.DataFrame, cols: list[str], stats_map: dict | None = None):
    raw = pd.DataFrame(index=df.index)
    for col in cols:
        x = pd.to_numeric(df[col], errors="coerce").astype(float)
        if col in LOG_TRANSFORM:
            x = np.log1p(x.clip(lower=0))
        raw[col] = x

    if stats_map is None:
        stats_map = {}
        fit = True
    else:
        fit = False

    z = pd.DataFrame(index=df.index)
    for col in cols:
        if fit:
            mu = float(raw[col].mean())
            sd = float(raw[col].std(ddof=0))
            if not np.isfinite(sd) or sd <= 0:
                sd = 1.0
            stats_map[col] = (mu, sd)
        else:
            mu, sd = stats_map[col]
        z[col] = (raw[col] - mu) / sd
    return z, stats_map


def add_intercept(Xdf: pd.DataFrame):
    names = ["Intercept"] + list(Xdf.columns)
    X = np.column_stack([np.ones(len(Xdf)), Xdf.to_numpy(float)])
    return X, names


def ols_beta(X: np.ndarray, y: np.ndarray) -> np.ndarray:
    return np.linalg.lstsq(X, y, rcond=None)[0]


def ols_hc3(X: np.ndarray, y: np.ndarray):
    beta = ols_beta(X, y)
    resid = y - X @ beta
    xtx_inv = np.linalg.pinv(X.T @ X)
    h = np.sum((X @ xtx_inv) * X, axis=1)
    adj = resid / np.maximum(1.0 - h, 1e-8)
    meat = X.T @ ((adj[:, None] ** 2) * X)
    cov = xtx_inv @ meat @ xtx_inv
    return beta, cov


def effect_summary(effect: float, se: float, alt: str = ">") -> dict:
    z = effect / max(se, 1e-12)
    p_two = 2.0 * norm_sf(abs(z))
    p_one = norm_sf(z) if alt == ">" else norm_cdf(z)
    return {
        "estimate": effect,
        "se_delta": se,
        "z_delta": z,
        "p_delta_two_sided": p_two,
        "p_delta_one_sided": p_one,
        "ci_delta_low_95": effect - 1.96 * se,
        "ci_delta_high_95": effect + 1.96 * se,
    }


def bootstrap_effect(X: np.ndarray, y: np.ndarray, contrast: np.ndarray, n_boot=N_BOOT, seed=RNG_SEED):
    rng = np.random.default_rng(seed)
    n = len(y)
    vals = np.empty(n_boot)
    for b in range(n_boot):
        idx = rng.integers(0, n, size=n)
        vals[b] = float(contrast @ ols_beta(X[idx], y[idx]))
    return vals


def freedman_lane_permutation(
    X_full: np.ndarray,
    X_reduced: np.ndarray,
    y: np.ndarray,
    contrast: np.ndarray,
    obs: float,
    alt: str = ">",
    n_perm=N_PERM,
    seed=RNG_SEED + 17,
):
    rng = np.random.default_rng(seed)
    beta_red = ols_beta(X_reduced, y)
    fitted_red = X_reduced @ beta_red
    resid_red = y - fitted_red
    vals = np.empty(n_perm)
    for b in range(n_perm):
        y_perm = fitted_red + rng.permutation(resid_red)
        vals[b] = float(contrast @ ols_beta(X_full, y_perm))
    p_two = (np.sum(np.abs(vals) >= abs(obs)) + 1.0) / (n_perm + 1.0)
    if alt == ">":
        p_one = (np.sum(vals >= obs) + 1.0) / (n_perm + 1.0)
    else:
        p_one = (np.sum(vals <= obs) + 1.0) / (n_perm + 1.0)
    return vals, p_two, p_one


def summarize_boot(vals: np.ndarray, obs: float) -> dict:
    prob_gt0 = float(np.mean(vals > 0))
    prob_lt0 = float(np.mean(vals < 0))
    p_sign_two = 2 * min(prob_gt0, prob_lt0)
    return {
        "bootstrap_mean": float(np.mean(vals)),
        "bootstrap_median": float(np.median(vals)),
        "bootstrap_ci_low_95": float(np.percentile(vals, 2.5)),
        "bootstrap_ci_high_95": float(np.percentile(vals, 97.5)),
        "bootstrap_prob_gt0": prob_gt0,
        "bootstrap_sign_p_two_sided": min(1.0, p_sign_two),
        "bootstrap_sign_p_one_sided_gt0": prob_lt0,
    }


def high_rt_stratified(df: pd.DataFrame) -> dict:
    groups = pd.qcut(df["Res_time"].rank(method="first"), q=3, labels=["Low", "Medium", "High"])
    sub = df.loc[groups == "High"].copy()
    Xz, _ = transform_matrix(sub, PREDICTORS)
    X, names = add_intercept(Xz)
    y = sub["ln"].to_numpy(float)
    beta, cov = ols_hc3(X, y)
    fish_idx = names.index("fish_gdp_sqkm")
    contrast = np.zeros(len(names))
    contrast[fish_idx] = 1.0
    obs = float(contrast @ beta)
    se = math.sqrt(max(float(contrast @ cov @ contrast), 0.0))
    summ = effect_summary(obs, se, alt=">")

    reduced_cols = [c for c in PREDICTORS if c != "fish_gdp_sqkm"]
    Xred_z = Xz[reduced_cols]
    Xred, _ = add_intercept(Xred_z)
    boot = bootstrap_effect(X, y, contrast, seed=RNG_SEED + 1)
    perm, p_perm_two, p_perm_one = freedman_lane_permutation(
        X, Xred, y, contrast, obs, alt=">", seed=RNG_SEED + 2
    )
    out = {
        "analysis": "High-RT stratified GLM fishery slope",
        "N": len(sub),
        "RT_days": "High tertile (>446.1 d)",
        "rt_z": np.nan,
        **summ,
        **summarize_boot(boot, obs),
        "permutation_p_two_sided": p_perm_two,
        "permutation_p_one_sided": p_perm_one,
        "test_note": "Pairs bootstrap and Freedman-Lane permutation within high-RT stratum.",
    }
    return out


def global_marginal_effects(df: pd.DataFrame) -> list[dict]:
    cols = PREDICTORS + ["Res_time"]
    Xz, stats_map = transform_matrix(df, cols)
    Xz["fish_gdp_sqkm:Res_time"] = Xz["fish_gdp_sqkm"] * Xz["Res_time"]
    X, names = add_intercept(Xz)
    y = df["ln"].to_numpy(float)
    beta, cov = ols_hc3(X, y)

    reduced_cols = [c for c in Xz.columns if c not in {"fish_gdp_sqkm", "fish_gdp_sqkm:Res_time"}]
    Xred, _ = add_intercept(Xz[reduced_cols])

    rt_vals = {
        "Upper-tertile boundary (446.1 d)": 446.1,
        "High-tertile median": float(df.loc[pd.qcut(df["Res_time"].rank(method="first"), q=3, labels=["Low", "Medium", "High"]) == "High", "Res_time"].median()),
        "RT 75th percentile": float(df["Res_time"].quantile(0.75)),
        "RT 90th percentile": float(df["Res_time"].quantile(0.90)),
    }
    mu, sd = stats_map["Res_time"]
    fish_idx = names.index("fish_gdp_sqkm")
    int_idx = names.index("fish_gdp_sqkm:Res_time")
    rows = []
    for label, rt in rt_vals.items():
        rt_z = (math.log1p(max(rt, 0.0)) - mu) / sd
        contrast = np.zeros(len(names))
        contrast[fish_idx] = 1.0
        contrast[int_idx] = rt_z
        obs = float(contrast @ beta)
        se = math.sqrt(max(float(contrast @ cov @ contrast), 0.0))
        summ = effect_summary(obs, se, alt=">")
        boot = bootstrap_effect(X, y, contrast, seed=RNG_SEED + int(round(rt)) % 10000)
        perm, p_perm_two, p_perm_one = freedman_lane_permutation(
            X, Xred, y, contrast, obs, alt=">", seed=RNG_SEED + 100 + int(round(rt)) % 10000
        )
        rows.append({
            "analysis": "Global GLM marginal fishery effect",
            "N": len(df),
            "RT_days": label,
            "rt_value_days": rt,
            "rt_z": rt_z,
            **summ,
            **summarize_boot(boot, obs),
            "permutation_p_two_sided": p_perm_two,
            "permutation_p_one_sided": p_perm_one,
            "test_note": "Marginal effect beta_fish + beta_interaction * RT_z from global interaction model.",
        })
    return rows


def write_docx(rows: pd.DataFrame) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    out = OUT / "Hydrology_marginal_bootstrap_permutation_precheck.docx"
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11)
    sec.page_height = Inches(8.5)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    sec.left_margin = sec.right_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    doc.styles["Normal"].font.size = Pt(12)

    def add_p(text, bold=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(12)

    def border(cell, top=None, bottom=None):
        tc_pr = cell._tc.get_or_add_tcPr()
        b = tc_pr.first_child_found_in("w:tcBorders")
        if b is None:
            b = OxmlElement("w:tcBorders")
            tc_pr.append(b)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = b.find(qn(f"w:{edge}"))
            if el is None:
                el = OxmlElement(f"w:{edge}")
                b.append(el)
            if edge == "top" and top:
                el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(top)); el.set(qn("w:color"), "000000")
            elif edge == "bottom" and bottom:
                el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(bottom)); el.set(qn("w:color"), "000000")
            else:
                el.set(qn("w:val"), "none")

    def cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(text))
        r.bold = bold
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(12)

    def fmtp(x):
        return "<0.001" if x < 0.001 else f"{x:.3f}"

    add_p("Hydrology preliminary significance checks for long-residence fishery effect", bold=True)
    headers = ["Analysis", "RT condition", "N", "Estimate", "Delta 95% CI", "Delta p", "Bootstrap 95% CI", "Perm. p"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, bold=True)
        border(table.rows[0].cells[i], top=12, bottom=4)
    for _, r in rows.iterrows():
        cells = table.add_row().cells
        vals = [
            r["analysis"],
            r["RT_days"],
            int(r["N"]),
            f"{r['estimate']:.3f}",
            f"{r['ci_delta_low_95']:.3f} to {r['ci_delta_high_95']:.3f}",
            fmtp(r["p_delta_two_sided"]),
            f"{r['bootstrap_ci_low_95']:.3f} to {r['bootstrap_ci_high_95']:.3f}",
            fmtp(r["permutation_p_two_sided"]),
        ]
        for i, val in enumerate(vals):
            cell_text(cells[i], val, align=WD_ALIGN_PARAGRAPH.LEFT if i in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER)
            border(cells[i])
    for c in table.rows[-1].cells:
        border(c, bottom=12)
    add_p("Note: Delta p and permutation p shown here are two-sided. One-sided values are retained in the CSV for directional-hypothesis reporting.")
    doc.save(out)
    return out


def main():
    df = load_data()
    rows = [high_rt_stratified(df)] + global_marginal_effects(df)
    out = pd.DataFrame(rows)
    csv = OUT / "Hydrology_marginal_bootstrap_permutation_precheck.csv"
    out.to_csv(csv, index=False, encoding="utf-8-sig")
    docx = write_docx(out)
    print(f"Saved CSV: {csv}")
    print(f"Saved DOCX: {docx}")
    show_cols = [
        "analysis", "RT_days", "N", "estimate", "se_delta", "p_delta_two_sided",
        "bootstrap_ci_low_95", "bootstrap_ci_high_95", "permutation_p_two_sided",
        "p_delta_one_sided", "permutation_p_one_sided",
    ]
    print(out[show_cols].to_string(index=False))


if __name__ == "__main__":
    main()
